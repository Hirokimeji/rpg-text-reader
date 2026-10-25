import hashlib
import json
import struct
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, MagicMock

from engine import extract, game_root
from exporter import export_records
from srpg import clean_srpg, emit_document, project_data, run_parser, parser_path


def archive(payload=b'\0' * 16, offset=0, version=1306, encrypted=0):
    return struct.pack('<6I', 0x53544453, encrypted, version, 1, 0, offset) + b'\0' * 144 + b'x' * offset + payload


def demo_map():
    def event(name):
        return {'id': 0, 'name': name, 'desc': '', 'pages': [
            {'id': 0, 'commands': [
                {'id': 1, 'type': 'message', 'data': ['你好。\n出发吧。'], 'speaker': '小明'},
                {'id': 2, 'type': 'choice', 'data': ['出发', '再等等']},
                {'id': 3, 'type': 'message', 'data': ['旁白'], 'speaker': 'No Speaker'},
                {'id': 4, 'type': 'message', 'data': ['未知角色'], 'speaker': 'Unknown_4294967295'},
                {'id': 5, 'type': 'script', 'data': ['do_not_execute()']},
                {'id': 6, 'infoText': '获得物品', 'comment': '编辑备注'},
            ]},
            {'id': 1, 'commands': [{'id': 1, 'type': 'message', 'data': ['你好。\n出发吧。'], 'speaker': '小明'}]},
        ]}
    return {'id': 2, 'name': '测试地图', 'mapName': '森林', 'victoryConds': ['抵达出口'],
            'openingEvents': [event('开始')], 'autoEvents': [event('自动')],
            'EnemyUnits': [{'id': 4, 'name': '守卫', 'events': [event('交谈')]}]}


def write_demo(project, output, tick):
    output.mkdir()
    (output / 'Maps').mkdir()
    (output / 'Maps/map_002.json').write_text(json.dumps(demo_map(), ensure_ascii=False), encoding='utf-8')
    (output / 'titles.json').write_text('{"gameTitle": "示例游戏"}', encoding='utf-8')
    tick('示例解析完成')


class SrpgTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.data = self.root / 'data.dts'
        self.data.write_bytes(archive())
        (self.root / 'Game.exe').write_bytes(b'never executed')

    def tearDown(self):
        self.tmp.cleanup()

    def test_detection_and_project_segment(self):
        self.data.write_bytes(archive(b'project text here', offset=20))
        for source in (self.root, self.data, self.root / 'Game.exe'):
            self.assertEqual(game_root(source), self.root.resolve())
        version, segments, payload = project_data(self.data)
        self.assertEqual((version, segments, payload), (1306, 0, b'project text here'))

    def test_invalid_unsupported_and_encrypted_data(self):
        for content, error in (
            (b'SDTS', '文件头'), (archive(version=1323), '数据版本'),
            (archive(encrypted=1), '已加密'),
            (archive()[:20] + struct.pack('<I', 99999) + archive()[24:], '越界'),
        ):
            with self.subTest(error=error):
                self.data.write_bytes(content)
                with self.assertRaisesRegex(ValueError, error):
                    project_data(self.data)

    def test_scene_identity_dialogue_choices_and_metadata(self):
        with patch('srpg.run_parser', side_effect=write_demo):
            result = extract(str(self.root / 'Game.exe'))
        self.assertEqual(result.engine, 'SRPG Studio')
        self.assertEqual(result.title, '示例游戏')
        repeated = [r for r in result.records if r.raw == '你好。\n出发吧。']
        self.assertEqual(len(repeated), 6)
        self.assertEqual(len({r.scene_id for r in repeated}), 6)
        self.assertEqual(len({r.source for r in result.records}), len(result.records))
        self.assertEqual(sum(r.category == '选项' for r in result.records), 6)
        self.assertTrue(all(not r.speaker for r in result.records if r.raw in ('旁白', '未知角色')))
        self.assertTrue(any('敌方单位' in r.scene and '守卫' in r.scene for r in repeated))
        self.assertTrue(any(r.raw == 'do_not_execute()' and r.category == '脚本／插件（待辨识）' for r in result.records))
        self.assertFalse(any(r.raw in ('message', 'No Speaker') for r in result.records))

    def test_help_pages_are_not_event_pages(self):
        rows = []
        emit_document([{'id': 2, 'name': '帮助', 'pages': ['第一页说明', '第二页说明']}],
                      'Extra/help.json', lambda *args: rows.append(args), lambda _: None)
        self.assertEqual([r[0] for r in rows], ['帮助', '第一页说明', '第二页说明'])

    def test_controls_keep_srpg_variables(self):
        self.assertEqual(clean_srpg(r'\C[2]\F[1]\fs[24]你好\wa[30]\_\vo[line.ogg]\va1[2]\p[3]\Mystery[4]'),
                         r'你好\va1[2]\p[3]\Mystery[4]')

    def test_translation_and_exports_preserve_source(self):
        translation = self.root / 'translation.json'
        translation.write_text('{"你好。\\n出发吧。": "Hello.\\nLet us go.", "小明": "Ming"}', encoding='utf-8')
        before = hashlib.sha256(self.data.read_bytes()).hexdigest()
        with patch('srpg.run_parser', side_effect=write_demo):
            result = extract(str(self.root), str(translation))
        dialogue = next(r for r in result.records if r.speaker == 'Ming')
        self.assertEqual(dialogue.translation, 'Hello.\nLet us go.')
        txt = export_records(result, result.records, str(self.root / 'out'), split=True)
        self.assertEqual(len(list(txt.glob('*.txt'))), result.stats['scenes'])
        xlsx = export_records(result, result.records, str(self.root / 'out'), fmt='xlsx')
        from openpyxl import load_workbook
        book = load_workbook(next(xlsx.glob('*.xlsx')), read_only=True)
        try:
            self.assertEqual(sum(1 for _ in book['文本'].iter_rows(min_row=2)), len(result.records))
        finally:
            book.close()
        self.assertEqual(before, hashlib.sha256(self.data.read_bytes()).hexdigest())

    def test_cancel_stops_helper_process(self):
        process = MagicMock()
        process.poll.return_value = None
        with patch('srpg.subprocess.Popen', return_value=process):
            with self.assertRaises(InterruptedError):
                run_parser(self.root / 'project.dat', self.root / 'out',
                           lambda _: (_ for _ in ()).throw(InterruptedError('已取消')))
        process.kill.assert_called_once()
        process.wait.assert_called()

    def test_native_helper_rejects_invalid_project(self):
        # Exercises the shipped helper and its error handling without game data.
        parser_path()
        (self.root / 'config.json').write_text('{"fileVersion": 1306, "segments": 0}')
        project = self.root / 'project.dat'
        project.write_bytes(b'\0' * 16)
        with self.assertRaisesRegex(ValueError, '解析失败'):
            run_parser(project, self.root / 'text', lambda _: None, timeout=5)


if __name__ == '__main__':
    unittest.main()
