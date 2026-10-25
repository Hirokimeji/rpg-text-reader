import json
import struct
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from engine import extract, game_root
from exporter import export_records
from unity_reader import ManagedReader, clean_unity, decode_text, emit_text_tree, emit_text_asset, extract_unity, unity_data


def demo_dialogue():
    return {'m_Name': 'Dialogue_01', 'DialogueName': '林间相遇', 'dialogue': {'DialogueLine': [
        {'Index': 0, 'Name': '旅行者', 'Text': '<color=green>你好。</color>\n一起走吧。',
         'SEDescription': '演出备注', 'choices': [{'ChoiceText': '出发', 'ChoiceCondition': 'ready'}, {'ChoiceText': '等待'}]},
        {'Index': 1, 'Name': '', 'Text': '<耳语>风停了。', 'commandInfos': [{'Parameters': 'internal_asset_name'}]},
        {'Index': 2, 'Name': '旅行者', 'Text': '<color=green>你好。</color>\n一起走吧。'},
    ]}}


def demo_extract(root, add, tick):
    emit_text_tree(demo_dialogue(), 'Dialogue_01', 'unity/data.unity3d/resources.assets/7',
                   'data.unity3d/resources.assets:pathID=7', add, class_name='Demo.DialogueData', tick=tick)
    return 'Unity 示例', [], {'dialogue_objects': 1}


def node(kind, name, children=(), flags=0):
    from UnityPy.helpers.TypeTreeNode import TypeTreeNode
    n = TypeTreeNode(0, kind, name, 0, 0, m_MetaFlag=flags)
    n.m_Children = list(children)
    return n


def aligned_string(text):
    raw = text.encode('utf-8')
    return struct.pack('<i', len(raw)) + raw + b'\0' * (-len(raw) % 4)


class UnityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.data = self.root / 'Demo_Data'
        self.data.mkdir()
        (self.data / 'data.unity3d').write_bytes(b'fixture, never executed')
        self.exe = self.root / 'Demo.exe'
        self.exe.write_bytes(b'fixture')

    def tearDown(self):
        self.tmp.cleanup()

    def test_detection_folder_exe_and_data_folder(self):
        for p in (self.root, self.exe, self.data):
            self.assertEqual(game_root(p), self.root.resolve())
        other = self.root / 'Other_Data'
        other.mkdir()
        (other / 'globalgamemanagers').write_bytes(b'fixture')
        self.assertIsNone(unity_data(self.root))
        self.assertEqual(unity_data(self.data), self.data)

    def test_cleanup_preserves_semantic_brackets_variables_and_raw_source(self):
        self.assertEqual(clean_unity(r'<size=140%><b>你好</b></size>{pause:100}<br><念话>\V[3]'), '你好\n<念话>\\V[3]')
        self.assertEqual(decode_text('中文'.encode('utf-16')), '中文')
        with self.assertRaises(ValueError):
            decode_text(b'\x00binary')

    def test_dialogue_choices_duplicate_locations_and_exports(self):
        before = (self.data / 'data.unity3d').read_bytes()
        with patch('engine.extract_unity', side_effect=demo_extract):
            r = extract(str(self.exe))
        dialogue = [x for x in r.records if x.category == '对白']
        self.assertEqual(len(dialogue), 3)
        self.assertEqual(dialogue[0].original, '你好。\n一起走吧。')
        self.assertEqual(dialogue[0].raw, demo_dialogue()['dialogue']['DialogueLine'][0]['Text'])
        self.assertNotEqual(dialogue[0].source, dialogue[2].source)
        self.assertEqual(sum(x.category == '选项' for x in r.records), 2)
        self.assertEqual(next(x.context for x in r.records if x.original == '出发'), 'ready')
        self.assertFalse(any('internal_asset_name' in x.raw for x in r.records))
        self.assertTrue(any(x.category == '注释' and x.raw == '演出备注' for x in r.records))
        txt = export_records(r, r.records, str(self.root / 'out'))
        self.assertIn('旅行者：你好。', next(txt.glob('*.txt')).read_text(encoding='utf-8-sig'))
        xlsx = export_records(r, r.records, str(self.root / 'out'), fmt='xlsx')
        from openpyxl import load_workbook
        book = load_workbook(next(xlsx.glob('*.xlsx')), read_only=True)
        try:
            self.assertEqual(sum(1 for _ in book['文本'].iter_rows(min_row=2)), len(r.records))
        finally:
            book.close()
        self.assertEqual(before, (self.data / 'data.unity3d').read_bytes())

    def test_text_asset_json_csv_and_unknown_text(self):
        rows = []
        add = lambda *args: rows.append(args)
        emit_text_asset(b'{"greeting":"Hello","bye":"Goodbye"}', 'locale.json', 'json', 'locale', add, lambda _: None)
        self.assertEqual([x[0] for x in rows], ['Hello', 'Goodbye'])
        self.assertTrue(all(x[4] == '其他文本（待辨识）' for x in rows))
        rows.clear()
        emit_text_asset('speaker,text\n莉娜,"第一行\n第二行"\n'.encode(), 'dialogue.csv', 'csv', 'table', add, lambda _: None)
        self.assertEqual(rows[0][0], '第一行\n第二行')
        self.assertEqual(rows[0][6], '莉娜')
        rows.clear()
        emit_text_asset('独立的说明文字'.encode(), 'notes.txt', 'text', 'notes', add, lambda _: None)
        self.assertEqual(rows[0][4], '其他文本（待辨识）')

    def test_translation_applies_without_rpg_maker_control_substitution(self):
        f = self.root / 'translations.json'
        f.write_text(json.dumps({'<耳语>风停了。': r'Hello\V[3]', '旅行者': 'Traveler'}), encoding='utf-8')
        with patch('engine.extract_unity', side_effect=demo_extract):
            r = extract(str(self.exe), str(f))
        self.assertTrue(any(x.translation == r'Hello\V[3]' for x in r.records))
        self.assertTrue(any(x.speaker == 'Traveler' for x in r.records))

    def test_cancel_restores_unity_reader_configuration(self):
        from UnityPy.helpers import TypeTreeHelper
        old = TypeTreeHelper.read_typetree_boost
        def cancel(_):
            raise InterruptedError('已取消')
        with self.assertRaises(InterruptedError):
            extract_unity(self.root, lambda *args: None, cancel)
        self.assertIs(TypeTreeHelper.read_typetree_boost, old)

    def test_generated_layout_clone_and_payload_header(self):
        root = node('MonoBehaviour', 'Base', [node('string', 'm_Name'), node('string', 'Text')])
        fake = SimpleNamespace(get_nodes_up=lambda *args: root)
        reader = ManagedReader(self.data)
        reader.generators['6000.0'] = fake
        payload = reader.node('6000.0', 'Demo.Line', 'Demo', payload=True)
        self.assertEqual([x.m_Name for x in payload.m_Children], ['Text'])
        self.assertEqual([x.m_Name for x in root.m_Children], ['m_Name', 'Text'])

    def test_managed_reference_registry_is_validated_byte_for_byte(self):
        from UnityPy.streams import EndianBinaryReader
        ref = node('Ref', 'data', [node('SInt64', 'rid')])
        vector = node('vector', 'eventDatas', [node('Array', 'Array', [node('int', 'size'), ref])])
        root = node('MonoBehaviour', 'Base', [vector, node('ManagedReferencesRegistry', 'references')])
        payload = node('Demo.Line', 'Base', [node('string', 'Text')])
        raw = struct.pack('<iqiiq', 1, 7, 2, 1, 7) + aligned_string('Line') + aligned_string('Demo') + aligned_string('Demo') + aligned_string('测试台词')
        def obj(data):
            reader = EndianBinaryReader(data, endian='<')
            return SimpleNamespace(serialized_type=SimpleNamespace(node=None),
                                   assets_file=SimpleNamespace(unity_version='6000.0'), reader=reader,
                                   byte_size=len(data), reset=lambda: setattr(reader, 'Position', 0))
        m = ManagedReader(self.data)
        with patch.object(m, 'node', side_effect=lambda *args, **kw: payload if kw.get('payload') else root):
            self.assertEqual(m.read(obj(raw), 'Demo.Event', 'Demo'), {'eventDatas': [{'Text': '测试台词'}]})
            with self.assertRaisesRegex(ValueError, '长度'):
                m.read(obj(raw + b'\0'), 'Demo.Event', 'Demo')


if __name__ == '__main__':
    unittest.main()
