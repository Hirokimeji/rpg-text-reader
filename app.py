from __future__ import annotations

import json
import os
import sys
import traceback
from collections import Counter, defaultdict
from pathlib import Path

from PySide6.QtCore import Qt, QAbstractTableModel, QModelIndex, QThread, Signal, QTimer, QUrl
from PySide6.QtGui import QColor, QDesktopServices, QFont, QIcon, QKeySequence, QShortcut
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QLineEdit, QFileDialog, QComboBox, QCheckBox, QSplitter, QTreeWidget, QTreeWidgetItem,
    QTableView, QHeaderView, QAbstractItemView, QTabWidget, QPlainTextEdit, QProgressBar,
    QMessageBox, QDialog, QDialogButtonBox, QFormLayout, QSpinBox,
)

from engine import STORY, Result, extract, game_root, translation_candidates
from exporter import export_records, record_text


APP_NAME = '台词拾光'
APP_VERSION = '1.1.1'
STYLE = '''
QWidget { font-family: "Microsoft YaHei UI"; font-size: 13px; color: #253347; }
QMainWindow, QDialog { background: #f3f6fa; }
QLabel#brand { font-size: 25px; font-weight: 700; color: #143a46; }
QLabel#muted { color: #65758b; }
QLabel#section { font-size: 15px; font-weight: 600; }
QLabel#badge { color: #146c67; background: #e1f1ec; border-radius: 8px; padding: 7px 12px; }
QLineEdit, QComboBox, QSpinBox { background: white; border: 1px solid #d8e0e9; border-radius: 6px; padding: 7px; }
QLineEdit:focus, QComboBox:focus { border: 1px solid #27877e; }
QPushButton { background: white; border: 1px solid #d8e0e9; border-radius: 6px; padding: 8px 14px; }
QPushButton:hover { background: #eaf3f1; border-color: #7ea9a4; }
QPushButton#primary { background: #146c67; color: white; border: none; font-weight: 600; }
QPushButton#primary:hover { background: #0b5754; }
QPushButton:disabled { color: #98a4b2; background: #e9edf2; }
QTreeWidget, QTableView, QPlainTextEdit { background: white; border: 1px solid #e0e6ed; border-radius: 7px; selection-background-color: #d8eeea; selection-color: #163b38; }
QTreeWidget::item { padding: 5px 2px; }
QTableView { gridline-color: #edf1f5; alternate-background-color: #f8fafc; }
QTableView::item { padding: 5px; }
QHeaderView::section { background: #edf2f6; color: #52657c; border: none; border-bottom: 1px solid #dde4ec; padding: 9px 6px; font-weight: 600; }
QTabWidget::pane { border: none; }
QTabBar::tab { padding: 9px 16px; background: #eaf0f5; border-top-left-radius: 5px; border-top-right-radius: 5px; }
QTabBar::tab:selected { background: white; color: #146c67; font-weight: 600; }
QSplitter::handle { background: #f3f6fa; }
QProgressBar { border: none; background: #e5ecef; border-radius: 3px; max-height: 5px; }
QProgressBar::chunk { background: #27877e; }
QCheckBox { spacing: 7px; }
QStatusBar { color: #65758b; }
'''


def label(text, name=''):
    widget = QLabel(text)
    if name:
        widget.setObjectName(name)
    return widget


def button(text, callback, primary=False):
    widget = QPushButton(text)
    widget.clicked.connect(callback)
    if primary:
        widget.setObjectName('primary')
    return widget


class Worker(QThread):
    progress = Signal(str)
    succeeded = Signal(object)
    failed = Signal(str)

    def __init__(self, work, parent=None):
        super().__init__(parent)
        self.work = work

    def run(self):
        try:
            value = self.work(self.progress.emit, self.isInterruptionRequested)
            self.succeeded.emit(value)
        except InterruptedError as exc:
            self.failed.emit(str(exc))
        except Exception as exc:
            log_error()
            self.failed.emit(f'{type(exc).__name__}: {exc}')
        finally:
            self.work = None


class RecordsModel(QAbstractTableModel):
    headers = ['序号', '类别', '说话人', '文本预览', '匹配状态']

    def __init__(self, parent=None):
        super().__init__(parent)
        self.records = []
        self.indices = []
        self.original = False

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.indices)

    def columnCount(self, parent=QModelIndex()):
        return len(self.headers)

    def record(self, row):
        return self.records[self.indices[row]]

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        r = self.record(index.row())
        if role == Qt.ItemDataRole.DisplayRole:
            return (str(r.seq), r.category, r.speaker or '—',
                    (r.original if self.original else r.display).replace('\n', '  ')[:500], r.status)[index.column()]
        if role == Qt.ItemDataRole.ToolTipRole:
            return f'{r.group} / {r.scene}\n{r.context}\n{r.display[:2000]}'
        if role == Qt.ItemDataRole.ForegroundRole and index.column() == 4:
            return QColor('#a26b20' if r.status in ('未匹配', '部分匹配') else '#146c67')
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if orientation == Qt.Orientation.Horizontal and role == Qt.ItemDataRole.DisplayRole:
            return self.headers[section]
        return None

    def reset(self, records, indices, original=False):
        self.beginResetModel()
        self.records, self.indices, self.original = records, indices, original
        self.endResetModel()


class TranslationDialog(QDialog):
    def __init__(self, paths, parent=None):
        super().__init__(parent)
        self.setWindowTitle('选择文本来源')
        self.resize(680, 260)
        layout = QVBoxLayout(self)
        heading = label('检测到可能的翻译文件', 'section')
        layout.addWidget(heading)
        note = label('选择要关联到场景的译文版本。也可以只读取游戏内现有文字。', 'muted')
        note.setWordWrap(True)
        layout.addWidget(note)
        self.combo = QComboBox()
        for path in paths:
            self.combo.addItem(f'{path.name}  ·  {path.stat().st_size/1024/1024:.1f} MB', str(path))
        self.combo.addItem('只读取游戏内文本', '')
        layout.addWidget(self.combo)
        self.location = label('', 'muted')
        self.location.setWordWrap(True)
        self.combo.currentIndexChanged.connect(lambda: self.location.setText(self.combo.currentData()))
        self.location.setText(self.combo.currentData())
        layout.addWidget(self.location)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText('开始提取')
        buttons.button(QDialogButtonBox.StandardButton.Cancel).setText('取消')
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)


class ExportDialog(QDialog):
    def __init__(self, total, filtered, selected, parent=None):
        super().__init__(parent)
        self.setWindowTitle('导出参考文本')
        self.resize(600, 400)
        layout = QVBoxLayout(self)
        layout.addWidget(label('带走你需要的台词', 'brand'))
        layout.addWidget(label('每次导出都会建立新文件夹，保留之前的结果。', 'muted'))
        form = QFormLayout()
        self.scope = QComboBox()
        for text, key in [(f'当前筛选结果 · {filtered:,} 条', 'filtered'), (f'全部文本 · {total:,} 条', 'all'),
                          (f'选中的文本 · {selected:,} 条', 'selected')]:
            self.scope.addItem(text, key)
        self.format = QComboBox()
        self.format.addItem('TXT · 方便阅读和复制', 'txt')
        self.format.addItem('Excel (.xlsx) · 保留原文、译文与来源', 'xlsx')
        self.split = QCheckBox('每个事件页／公共事件分别生成一个文件')
        self.bilingual = QCheckBox('TXT 同时包含原文和译文')
        self.raw = QCheckBox('TXT 保留原始控制符')
        self.format.currentIndexChanged.connect(lambda: [w.setEnabled(self.format.currentData() == 'txt') for w in (self.bilingual, self.raw)])
        self.folder = QLineEdit(str(Path.home() / 'Documents' / '台词拾光导出'))
        path_row = QHBoxLayout()
        path_row.addWidget(self.folder)
        path_row.addWidget(button('选择…', self.browse))
        form.addRow('导出范围', self.scope)
        form.addRow('文件格式', self.format)
        form.addRow('场景拆分', self.split)
        form.addRow('阅读选项', self.bilingual)
        form.addRow('', self.raw)
        form.addRow('保存位置', path_row)
        layout.addLayout(form)
        note = label('Excel 始终包含原始文本；重复台词按各自场景保留。', 'muted')
        note.setWordWrap(True)
        layout.addWidget(note)
        controls = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        controls.button(QDialogButtonBox.StandardButton.Ok).setText('导出')
        controls.button(QDialogButtonBox.StandardButton.Cancel).setText('取消')
        controls.accepted.connect(self.accept)
        controls.rejected.connect(self.reject)
        layout.addWidget(controls)

    def browse(self):
        path = QFileDialog.getExistingDirectory(self, '选择导出位置', self.folder.text())
        if path:
            self.folder.setText(path)


class ReadingDialog(QDialog):
    """Pagination prevents a whole game's text from blocking the rich text widget."""
    def __init__(self, records, parent=None):
        super().__init__(parent)
        self.setWindowTitle('连续阅读 · 当前筛选结果')
        self.resize(850, 720)
        self.records, self.page = records, 0
        layout = QVBoxLayout(self)
        self.info = label('', 'section')
        layout.addWidget(self.info)
        self.body = QPlainTextEdit()
        self.body.setReadOnly(True)
        self.body.setStyleSheet('font-size: 18px; padding: 18px;')
        layout.addWidget(self.body)
        nav = QHBoxLayout()
        self.prev = button('上一页', lambda: self.turn(-1))
        self.next = button('下一页', lambda: self.turn(1))
        nav.addWidget(self.prev)
        nav.addStretch()
        nav.addWidget(button('复制本页', lambda: QApplication.clipboard().setText(self.body.toPlainText())))
        nav.addWidget(self.next)
        layout.addLayout(nav)
        self.turn(0)

    def turn(self, delta):
        self.page += delta
        start = self.page * 100
        subset = self.records[start:start+100]
        blocks, last = [], None
        for r in subset:
            if last != r.scene_id:
                blocks.append(f'【{r.group} / {r.scene}】\n')
                last = r.scene_id
            blocks.append(record_text(r))
        self.body.setPlainText('\n\n'.join(blocks))
        self.info.setText(f'第 {self.page+1} 页 · {start+1:,}–{start+len(subset):,} / {len(self.records):,} 条')
        self.prev.setEnabled(self.page > 0)
        self.next.setEnabled(start+100 < len(self.records))


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.result: Result | None = None
        self.worker = None
        self.filtered_indices = []
        self.table_page = 0
        self.scene_filter = ('all', '')
        self.setWindowTitle(f'{APP_NAME} {APP_VERSION} · RPG 游戏文本提取器')
        self.resize(1440, 900)
        self.setMinimumSize(1050, 680)
        self.setAcceptDrops(True)
        base = QWidget()
        self.setCentralWidget(base)
        layout = QVBoxLayout(base)
        layout.setContentsMargins(24, 18, 24, 12)
        layout.setSpacing(12)

        top = QHBoxLayout()
        names = QVBoxLayout()
        names.addWidget(label(APP_NAME, 'brand'))
        names.addWidget(label('从游戏对白，找到创作的线索', 'muted'))
        top.addLayout(names)
        top.addStretch()
        self.badge = label('本地读取 · MV / MZ / VX Ace', 'badge')
        top.addWidget(self.badge)
        self.report_button = button('提取报告', self.report)
        self.report_button.setEnabled(False)
        top.addWidget(self.report_button)
        top.addWidget(button('使用帮助', self.help))
        self.export_button = button('导出文本', self.export, True)
        self.export_button.setEnabled(False)
        top.addWidget(self.export_button)
        layout.addLayout(top)

        source = QHBoxLayout()
        self.path = QLineEdit()
        self.path.setPlaceholderText('选择游戏文件夹，或把游戏 EXE / 文件夹 / 翻译 JSON 拖到窗口')
        self.path.returnPressed.connect(lambda: self.load(self.path.text()))
        source.addWidget(self.path, 1)
        self.open_folder = button('选择游戏文件夹', self.browse_folder, True)
        self.open_file = button('选择 EXE / JSON', self.browse_file)
        self.translation_button = button('切换翻译文件', self.choose_translation)
        source.addWidget(self.open_folder)
        source.addWidget(self.open_file)
        source.addWidget(self.translation_button)
        layout.addLayout(source)

        filters = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText('搜索台词、原文、说话人或场景…  Ctrl+F')
        self.search.setClearButtonEnabled(True)
        filters.addWidget(self.search, 1)
        self.category = QComboBox()
        self.category.setMinimumWidth(170)
        self.category.addItem('剧情对白与选项', 'story')
        self.category.addItem('所有文本类别', 'all')
        filters.addWidget(self.category)
        self.status_filter = QComboBox()
        self.status_filter.addItems(['全部匹配状态', '已匹配', '部分匹配', '未匹配', '游戏内文本'])
        filters.addWidget(self.status_filter)
        self.view = QComboBox()
        self.view.addItems(['译文优先', '原文优先'])
        filters.addWidget(self.view)
        filters.addWidget(button('重置筛选', self.reset_filters))
        layout.addLayout(filters)

        self.summary = label('还没有打开游戏。选择文件夹即可开始，无需启动游戏或 MTool。', 'muted')
        self.summary.setWordWrap(True)
        layout.addWidget(self.summary)
        splitter = QSplitter()
        splitter.setChildrenCollapsible(False)
        left = QWidget()
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(0, 0, 0, 0)
        left_layout.addWidget(label('场景目录', 'section'))
        self.tree = QTreeWidget()
        self.tree.setHeaderHidden(True)
        self.tree.setMinimumWidth(200)
        self.tree.setUniformRowHeights(True)
        self.tree.itemSelectionChanged.connect(self.scene_changed)
        left_layout.addWidget(self.tree)
        left_layout.addWidget(label('地图 → 事件 → 事件页\n场景数量包含所有文本类别', 'muted'))
        splitter.addWidget(left)

        center = QWidget()
        center_layout = QVBoxLayout(center)
        center_layout.setContentsMargins(0, 0, 0, 0)
        center_top = QHBoxLayout()
        self.count = label('文本列表', 'section')
        center_top.addWidget(self.count)
        center_top.addStretch()
        self.read_button = button('连续阅读', self.read)
        self.read_button.setEnabled(False)
        center_top.addWidget(self.read_button)
        center_layout.addLayout(center_top)
        self.model = RecordsModel(self)
        self.table = QTableView()
        self.table.setModel(self.model)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setWordWrap(False)
        self.table.verticalHeader().hide()
        self.table.verticalHeader().setDefaultSectionSize(38)
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Stretch)
        for col, width in ((0, 64), (1, 80), (2, 85), (4, 95)):
            self.table.setColumnWidth(col, width)
        self.table.setMinimumWidth(380)
        self.table.selectionModel().selectionChanged.connect(self.show_record)
        center_layout.addWidget(self.table)
        pagination = QHBoxLayout()
        self.prev_page = button('上一页', lambda: self.turn_table(-1))
        self.next_page = button('下一页', lambda: self.turn_table(1))
        self.page_label = label('每页 200 条', 'muted')
        pagination.addWidget(self.prev_page)
        pagination.addWidget(self.page_label, 1, Qt.AlignmentFlag.AlignCenter)
        pagination.addWidget(self.next_page)
        self.prev_page.setEnabled(False)
        self.next_page.setEnabled(False)
        center_layout.addLayout(pagination)
        self.copy_results_button = button('复制当前筛选结果（可跨页）', self.copy_results)
        self.copy_results_button.setEnabled(False)
        center_layout.addWidget(self.copy_results_button)
        center_layout.addWidget(label('点击阅读详情 · Ctrl / Shift 多选 · Ctrl+C 复制选中台词', 'muted'))
        splitter.addWidget(center)

        right = QWidget()
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(0, 0, 0, 0)
        detail_top = QHBoxLayout()
        detail_top.addWidget(label('台词阅读', 'section'))
        detail_top.addStretch()
        detail_top.addWidget(label('字号', 'muted'))
        self.font_size = QSpinBox()
        self.font_size.setRange(12, 28)
        self.font_size.setValue(17)
        detail_top.addWidget(self.font_size)
        right_layout.addLayout(detail_top)
        self.detail_title = label('选中一段文字，开始阅读', 'muted')
        self.detail_title.setWordWrap(True)
        self.detail_title.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        right_layout.addWidget(self.detail_title)
        self.tabs = QTabWidget()
        self.panes = []
        for name in ('阅读', '原文', '译文', '原始／来源'):
            pane = QPlainTextEdit()
            pane.setReadOnly(True)
            pane.setPlaceholderText('提取完成后，在左侧选择场景或在中间选择台词。')
            self.tabs.addTab(pane, name)
            self.panes.append(pane)
        self.set_read_font()
        self.font_size.valueChanged.connect(self.set_read_font)
        right_layout.addWidget(self.tabs)
        self.copy_button = button('复制选中台词', self.copy)
        self.copy_button.setEnabled(False)
        right_layout.addWidget(self.copy_button)
        right.setMinimumWidth(300)
        splitter.addWidget(right)
        splitter.setSizes([250, 650, 440])
        layout.addWidget(splitter, 1)

        bottom = QHBoxLayout()
        self.activity = label('就绪 · 只读游戏文件', 'muted')
        bottom.addWidget(self.activity, 1)
        self.cancel_button = button('取消', self.cancel)
        self.cancel_button.hide()
        bottom.addWidget(self.cancel_button)
        layout.addLayout(bottom)
        self.progress = QProgressBar()
        self.progress.setRange(0, 0)
        self.progress.hide()
        layout.addWidget(self.progress)

        self.search_timer = QTimer(self)
        self.search_timer.setSingleShot(True)
        self.search_timer.setInterval(220)
        self.search_timer.timeout.connect(self.filter)
        self.search.textChanged.connect(lambda: self.search_timer.start())
        for combo in (self.category, self.status_filter, self.view):
            combo.currentIndexChanged.connect(self.filter)
        QShortcut(QKeySequence('Ctrl+F'), self, activated=self.search.setFocus)
        QShortcut(QKeySequence('Ctrl+O'), self, activated=self.browse_folder)
        copy_shortcut = QShortcut(QKeySequence('Ctrl+C'), self.table, activated=self.copy)
        copy_shortcut.setContext(Qt.ShortcutContext.WidgetShortcut)

    def set_read_font(self):
        for pane in self.panes:
            pane.setStyleSheet(f'font-size: {self.font_size.value()}px; padding: 12px;')

    def browse_folder(self):
        path = QFileDialog.getExistingDirectory(self, '选择游戏主文件夹')
        if path:
            self.load(path)

    def browse_file(self):
        path, _ = QFileDialog.getOpenFileName(self, '选择游戏 EXE 或翻译文件', '', '游戏或翻译 (*.exe *.json);;所有文件 (*)')
        if path:
            self.load(path)

    def choose_translation(self):
        if self.busy():
            return
        path, _ = QFileDialog.getOpenFileName(self, '选择 MTool 翻译 JSON', '', '翻译文件 (*.json)')
        if path:
            root = self.result.root if self.result and game_root(self.result.root) else path
            self.load(root, path)

    def load(self, path, translation=None):
        if self.busy():
            return
        path = str(path).strip().strip('"')
        if not path or not Path(path).exists():
            QMessageBox.warning(self, '路径不可用', '请选择存在的游戏文件夹、游戏 EXE 或翻译 JSON。')
            return
        if Path(path).suffix.lower() == '.json' and Path(path).is_file():
            translation = path
        root = game_root(path)
        if translation is None and root:
            candidates = translation_candidates(root)
            if candidates:
                dialog = TranslationDialog(candidates, self)
                if dialog.exec() != QDialog.DialogCode.Accepted:
                    return
                translation = dialog.combo.currentData()
        self.path.setText(path)
        self.start_worker(lambda progress, cancelled: extract(path, translation or '', progress, cancelled), self.loaded)

    def busy(self):
        return bool(self.worker and self.worker.isRunning())

    def start_worker(self, work, success):
        self.worker = Worker(work, self)
        self.worker.progress.connect(self.activity.setText)
        self.worker.succeeded.connect(success)
        self.worker.failed.connect(self.failed)
        self.worker.finished.connect(self.finished)
        for widget in (self.open_folder, self.open_file, self.translation_button, self.path, self.export_button):
            widget.setEnabled(False)
        self.progress.show()
        self.cancel_button.show()
        self.cancel_button.setEnabled(True)
        self.activity.setText('正在准备…')
        self.worker.start()

    def cancel(self):
        if self.busy():
            self.worker.requestInterruption()
            self.cancel_button.setEnabled(False)
            self.activity.setText('正在取消，请稍候…')

    def failed(self, text):
        self.activity.setText(text)
        if not text.startswith('已取消'):
            QMessageBox.warning(self, '操作未完成', text + '\n\n游戏文件未被修改。')

    def finished(self):
        self.progress.hide()
        self.cancel_button.hide()
        for widget in (self.open_folder, self.open_file, self.translation_button, self.path):
            widget.setEnabled(True)
        self.export_button.setEnabled(bool(self.result and self.result.records))
        if self.worker:
            self.worker.deleteLater()
            self.worker = None

    def loaded(self, result):
        self.result = result
        self.badge.setText(result.engine + ' · 只读')
        self.setWindowTitle(f'{APP_NAME} {APP_VERSION} · {result.title}')
        self.search.blockSignals(True)
        self.search.clear()
        self.search.blockSignals(False)
        self.category.blockSignals(True)
        self.category.clear()
        self.category.addItem('剧情对白与选项', 'story')
        self.category.addItem('所有文本类别', 'all')
        for category, count in result.stats['categories'].items():
            self.category.addItem(f'{category} · {count:,}', category)
        if not any(r.category in STORY for r in result.records):
            self.category.setCurrentIndex(1)
        self.category.blockSignals(False)
        self.status_filter.blockSignals(True)
        self.status_filter.setCurrentIndex(0)
        self.status_filter.blockSignals(False)
        self.build_tree()
        self.report_button.setEnabled(True)
        self.activity.setText(f'提取完成 · {len(result.records):,} 条 · 已保留来源位置')
        self.filter()

    def build_tree(self):
        self.tree.blockSignals(True)
        self.tree.clear()
        all_item = QTreeWidgetItem(['全部场景'])
        all_item.setData(0, Qt.ItemDataRole.UserRole, ('all', ''))
        self.tree.addTopLevelItem(all_item)
        groups = defaultdict(dict)
        counts = Counter(r.scene_id for r in self.result.records)
        for r in self.result.records:
            groups[r.group].setdefault(r.scene_id, r.scene)
        for group, scenes in groups.items():
            group_item = QTreeWidgetItem([f'{group}  ({len(scenes)})'])
            group_item.setData(0, Qt.ItemDataRole.UserRole, ('group', group))
            group_item.setToolTip(0, group)
            all_item.addChild(group_item)
            events = {}
            for sid, scene in scenes.items():
                parent = group_item
                title = scene
                if '/event' in sid and '/page' in sid:
                    eid = sid.rsplit('/page', 1)[0]
                    if eid not in events:
                        event_item = QTreeWidgetItem([scene.rsplit(' / 第 ', 1)[0]])
                        event_item.setData(0, Qt.ItemDataRole.UserRole, ('prefix', eid + '/'))
                        group_item.addChild(event_item)
                        events[eid] = event_item
                    parent = events[eid]
                    title = '第 ' + scene.rsplit(' / 第 ', 1)[-1]
                item = QTreeWidgetItem([f'{title}  ({counts[sid]:,})'])
                item.setData(0, Qt.ItemDataRole.UserRole, ('scene', sid))
                item.setToolTip(0, scene)
                parent.addChild(item)
        all_item.setExpanded(True)
        self.tree.setCurrentItem(all_item)
        self.scene_filter = ('all', '')
        self.tree.blockSignals(False)

    def scene_changed(self):
        current = self.tree.currentItem()
        self.scene_filter = current.data(0, Qt.ItemDataRole.UserRole) if current else ('all', '')
        self.filter()

    def filter(self):
        if not self.result:
            return
        query = self.search.text().casefold().strip()
        category = self.category.currentData()
        status = self.status_filter.currentText()
        mode, value = self.scene_filter
        indices = []
        for i, r in enumerate(self.result.records):
            if category == 'story' and r.category not in STORY:
                continue
            if category not in ('all', 'story') and r.category != category:
                continue
            if status != '全部匹配状态' and r.status != status:
                continue
            if mode == 'group' and r.group != value:
                continue
            if mode == 'scene' and r.scene_id != value:
                continue
            if mode == 'prefix' and not r.scene_id.startswith(value):
                continue
            if query and query not in r.search and query not in r.group.casefold():
                continue
            indices.append(i)
        self.filtered_indices = indices
        self.table_page = 0
        self.turn_table(0)
        self.count.setText(f'文本 · {len(indices):,} 条')
        source = Path(self.result.translation_file).name if self.result.translation_file else '游戏内现有文本'
        self.summary.setText(f'{self.result.title}  ·  共 {len(self.result.records):,} 条 / {self.result.stats["scenes"]:,} 个场景  ·  来源：{source}')
        self.read_button.setEnabled(bool(indices))
        self.copy_results_button.setEnabled(bool(indices))
        if not indices:
            self.clear_detail()
            self.detail_title.setText('没有符合条件的文本。可切换“所有文本类别”或重置筛选。')

    def turn_table(self, delta):
        if not self.result:
            return
        self.table_page = max(0, min(self.table_page + delta, max(0, (len(self.filtered_indices)-1)//200)))
        start = self.table_page * 200
        self.model.reset(self.result.records, self.filtered_indices[start:start+200], self.view.currentIndex() == 1)
        pages = max(1, (len(self.filtered_indices)+199)//200)
        self.page_label.setText(f'{self.table_page+1} / {pages:,} 页 · 每页 200 条')
        self.prev_page.setEnabled(self.table_page > 0)
        self.next_page.setEnabled(start + 200 < len(self.filtered_indices))
        if self.model.rowCount():
            self.table.selectRow(0)
            self.table.scrollToTop()

    def reset_filters(self):
        self.search.clear()
        self.category.setCurrentIndex(0)
        self.status_filter.setCurrentIndex(0)
        self.view.setCurrentIndex(0)
        if self.tree.topLevelItemCount():
            self.tree.setCurrentItem(self.tree.topLevelItem(0))
        self.filter()

    def selected_records(self):
        return [self.model.record(index.row()) for index in sorted(self.table.selectionModel().selectedRows(), key=lambda x: x.row())]

    def clear_detail(self):
        for pane in self.panes:
            pane.clear()
        self.copy_button.setEnabled(False)

    def show_record(self, *_):
        records = self.selected_records()
        if not records:
            self.clear_detail()
            return
        r = records[0]
        self.detail_title.setText(f'{r.group}\n{r.scene}\n{r.speaker or "说话人未标注"} · {r.category} · {r.status}' +
                                  (f' · 已选 {len(records)} 条' if len(records) > 1 else ''))
        self.panes[0].setPlainText((r.original if self.view.currentIndex() else r.display))
        self.panes[1].setPlainText(r.original)
        self.panes[2].setPlainText(r.translation or ('此处为游戏内现有文本，没有单独的原文／译文对照。' if r.status == '游戏内文本' else '未匹配到译文，阅读页显示原文。'))
        self.panes[3].setPlainText(f'来源位置\n{r.source}\n\n条件／分支\n{r.context or "无额外条件标记"}\n\n原始文本\n{r.raw}\n\n原始译文\n{r.raw_translation or "无"}')
        self.copy_button.setEnabled(True)

    def copy(self):
        records = self.selected_records()
        if records:
            QApplication.clipboard().setText('\n\n'.join(
                (f'{r.speaker}：' if r.speaker else '') + (r.original if self.view.currentIndex() else r.display) for r in records))
            self.activity.setText(f'已复制 {len(records):,} 条台词')

    def read(self):
        records = [self.result.records[i] for i in self.filtered_indices] if self.result else []
        if records:
            ReadingDialog(records, self).exec()

    def copy_results(self):
        if self.result and self.filtered_indices:
            records = [self.result.records[i] for i in self.filtered_indices]
            QApplication.clipboard().setText('\n\n'.join(record_text(r) for r in records))
            self.activity.setText(f'已复制当前筛选的 {len(records):,} 条文本（包括其他页）')

    def export(self):
        if not self.result or self.busy():
            return
        selected = self.selected_records()
        dialog = ExportDialog(len(self.result.records), len(self.filtered_indices), len(selected), self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        scope = dialog.scope.currentData()
        records = self.result.records if scope == 'all' else (selected if scope == 'selected' else [self.result.records[i] for i in self.filtered_indices])
        if not records:
            QMessageBox.information(self, '没有文本', '当前导出范围没有文本，请更换范围。')
            return
        target = dialog.folder.text().strip()
        if not target:
            QMessageBox.warning(self, '请选择位置', '导出位置不能为空。')
            return
        split = dialog.split.isChecked()
        scene_count = len({r.scene_id for r in records})
        if split and scene_count > 200:
            if QMessageBox.question(self, '场景文件较多', f'本次将生成 {scene_count:,} 个场景文件。是否继续？') != QMessageBox.StandardButton.Yes:
                return
        fmt, bilingual, raw = dialog.format.currentData(), dialog.bilingual.isChecked(), dialog.raw.isChecked()
        result = self.result
        self.start_worker(lambda progress, cancelled: export_records(result, records, target, fmt, split, bilingual, raw, progress, cancelled), self.exported)

    def exported(self, folder):
        self.activity.setText(f'已导出到 {folder}')
        msg = QMessageBox(self)
        msg.setWindowTitle('导出完成')
        msg.setText('文本已导出。')
        msg.setInformativeText(str(folder))
        open_button = msg.addButton('打开结果文件夹', QMessageBox.ButtonRole.AcceptRole)
        msg.addButton('关闭', QMessageBox.ButtonRole.RejectRole)
        msg.exec()
        if msg.clickedButton() == open_button:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(folder)))

    def report(self):
        if not self.result:
            return
        report = QDialog(self)
        report.setWindowTitle('提取报告')
        report.resize(760, 600)
        layout = QVBoxLayout(report)
        pane = QPlainTextEdit()
        pane.setReadOnly(True)
        r = self.result
        lines = [f'{r.title}', r.engine, f'游戏数据：{r.root}', f'翻译文件：{r.translation_file or "未使用"}', '',
                 f'文本记录：{len(r.records):,} 条', f'场景：{r.stats["scenes"]:,} 个',
                 f'翻译字典：{r.stats["translation_entries"]:,} 条', f'未关联场景的字典条目：{r.stats["unmatched_dictionary"]:,} 条',
                 '', '文本类别']
        lines.extend(f'  {key}：{value:,}' for key, value in r.stats['categories'].items())
        lines.extend(['', '匹配状态'])
        lines.extend(f'  {key}：{value:,}' for key, value in r.stats['statuses'].items())
        lines.extend(['', '提取说明', *('• ' + w for w in r.warnings)])
        pane.setPlainText('\n'.join(lines))
        layout.addWidget(pane)
        layout.addWidget(button('关闭', report.accept))
        report.exec()

    def help(self):
        QMessageBox.information(self, '开始使用台词拾光',
            '1. 选择游戏文件夹，或拖入游戏 EXE／翻译 JSON。\n'
            '2. 检测到多份翻译时，选择要使用的版本。\n'
            '3. 左侧选择场景，中间搜索台词，右侧查看原文、译文及来源。\n'
            '4. “连续阅读”按每页 100 条显示上下文；“导出文本”可汇总或按场景拆分。\n\n'
            '将本工具 EXE 放到 MV／MZ／VX Ace 游戏目录，双击会自动识别。\n'
            '默认只看对白和选项；物品、菜单、脚本等请切换“所有文本类别”。\n'
            '单独打开翻译 JSON 也能提取，但没有对应游戏数据时无法还原场景。\n\n'
            '〔变量#编号〕等表示游戏运行时决定的内容；未知插件控制符保留。\n'
            '“部分匹配”包含仍未翻译的行。“游戏内文本”不表示已自动判断语言。\n'
            '地图／事件页顺序不等于剧情先后顺序；公共事件单独保留，不自动重复展开。\n'
            '支持 MV／MZ 标准 JSON、VX Ace 已展开的 .rvdata2 与 MTool 字符串字典。\n'
            '暂不支持 XP／VX 的数据、RGSS 加密归档或特殊封装。\n'
            '提取与导出在本机完成，不启动游戏、不上传文本、不修改游戏。')

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls() and not self.busy():
            event.acceptProposedAction()

    def dropEvent(self, event):
        urls = event.mimeData().urls()
        if urls and urls[0].isLocalFile():
            self.load(urls[0].toLocalFile())

    def closeEvent(self, event):
        if self.busy():
            self.cancel()
            self.activity.setText('正在取消任务；结束后即可关闭窗口。')
            event.ignore()
        else:
            event.accept()


def log_error(exc_type=None, value=None, tb=None):
    folder = Path(os.environ.get('LOCALAPPDATA', str(Path.home()))) / '台词拾光'
    try:
        folder.mkdir(parents=True, exist_ok=True)
        with (folder / 'error.log').open('a', encoding='utf-8') as f:
            if exc_type:
                traceback.print_exception(exc_type, value, tb, file=f)
            else:
                traceback.print_exc(file=f)
    except OSError:
        pass


def main():
    # Offline packaging verification: never opens a window or accesses the clipboard.
    # This is a developer diagnostic path; ordinary launches always use the GUI.
    self_check = len(sys.argv) > 1 and sys.argv[1] == '--self-check'
    if self_check:
        if len(sys.argv) != 5:
            return 2
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setOrganizationName('Textlight')
    app.setStyle('Fusion')
    app.setStyleSheet(STYLE)
    if self_check:
        target = Path(sys.argv[4])
        target.mkdir(parents=True, exist_ok=True)
        try:
            result = extract(sys.argv[2], '' if sys.argv[3] == '-' else sys.argv[3])
            test_window = MainWindow()
            test_window.loaded(result)
            assert test_window.model.rowCount() <= 200
            assert len(test_window.filtered_indices) > 0
            subset = [result.records[i] for i in test_window.filtered_indices[:250]]
            txt = export_records(result, subset, str(target), 'txt')
            xlsx = export_records(result, subset, str(target), 'xlsx')
            from openpyxl import load_workbook
            workbook = load_workbook(next(xlsx.glob('*.xlsx')), read_only=True)
            row_ids = {int(row[0]) for row in workbook['文本'].iter_rows(min_row=2, values_only=True)}
            workbook.close()
            assert row_ids == {r.seq for r in subset}
            (target/'self-check.json').write_text(json.dumps({'ok':True, 'engine':result.engine, 'stats':result.stats,
                'txt':str(txt), 'xlsx':str(xlsx), 'exported_records':len(subset)}, ensure_ascii=False, indent=2), encoding='utf-8')
            test_window.close()
            return 0
        except Exception:
            (target/'self-check-error.txt').write_text(traceback.format_exc(), encoding='utf-8')
            return 1
    def exception_hook(kind, value, tb):
        log_error(kind, value, tb)
        QMessageBox.warning(None, '操作遇到问题', f'{kind.__name__}: {value}\n\n可重新打开游戏再试。错误日志位于本机 LocalAppData\\台词拾光\\error.log。')
    sys.excepthook = exception_hook
    icon = Path(getattr(sys, '_MEIPASS', Path(__file__).parent)) / 'app.ico'
    if icon.exists():
        app.setWindowIcon(QIcon(str(icon)))
    window = MainWindow()
    screen = app.primaryScreen().availableGeometry()
    window.resize(min(1440, screen.width()-60), min(900, screen.height()-60))
    window.show()
    if len(sys.argv) > 1:
        QTimer.singleShot(200, lambda: window.load(sys.argv[1]))
    else:
        folder = Path(sys.executable).parent if getattr(sys, 'frozen', False) else Path(__file__).parent
        if game_root(folder):
            QTimer.singleShot(200, lambda: window.load(str(folder)))
    return app.exec()


if __name__ == '__main__':
    sys.exit(main())
