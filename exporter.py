"""Exports are created in a new directory; source files are never overwritten."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from engine import Record, Result


def safe_name(name: str, limit=65):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).strip(' .')[:limit]
    if not name or name.upper() in {'CON', 'PRN', 'AUX', 'NUL', *(f'COM{i}' for i in range(10)), *(f'LPT{i}' for i in range(10))}:
        name = '_' + (name or '未命名')
    return name


def record_text(r: Record, bilingual=False, raw=False):
    text = (r.raw_translation or r.raw) if raw else r.display
    speaker = f'{r.speaker}：' if r.speaker else ''
    prefix = f'[{r.category}] ' if r.category != '对白' else ''
    status = f'〔{r.status}〕' if r.status in ('未匹配', '部分匹配') else ''
    body = prefix + status + speaker + text
    if bilingual and r.translation:
        body += '\n原文：' + (r.raw if raw else r.original)
    return body


def text_export(path, records, title, bilingual, raw, check):
    with path.open('w', encoding='utf-8-sig', newline='\n') as f:
        f.write(f'{title}\n台词拾光 · 游戏文本参考\n场景按数据结构排列，不代表实际游玩顺序。\n\n')
        previous = None
        for n, r in enumerate(records):
            if n % 500 == 0:
                check(f'正在导出 TXT · {n:,}/{len(records):,}')
            if previous != r.scene_id:
                f.write(f'\n{"="*48}\n{r.group} / {r.scene}\n{"="*48}\n\n')
                previous = r.scene_id
            f.write(record_text(r, bilingual, raw) + '\n\n')


def excel_export(path, records, result, check):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment

    book = Workbook(write_only=True)
    info = book.create_sheet('说明')
    info.column_dimensions['A'].width = 24
    info.column_dimensions['B'].width = 100
    def safe_cell(ws, value, heading=False):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))
        c = WriteOnlyCell(ws, value=value)
        c.data_type = 's'  # Do not execute formula-like game strings in Excel.
        c.alignment = Alignment(vertical='top', wrap_text=True)
        if heading:
            c.font = Font(color='FFFFFF', bold=True)
            c.fill = PatternFill('solid', fgColor='146C67')
        return c
    for key, value in [('游戏', result.title), ('引擎', result.engine), ('文本条数', len(records)),
                       ('翻译文件', result.translation_file or '未使用外部翻译'),
                       ('阅读顺序', '按地图、事件、事件页排列，不代表实际游玩顺序。'),
                       ('长文本', '超过 Excel 单元格长度的文字分为多个续段；序号和来源相同。'),
                       ('提取说明', '\n'.join(result.warnings))]:
        info.append([safe_cell(info, key), safe_cell(info, value)])
    headers = ['序号', '地图／分组', '场景', '类别', '说话人', '阅读文本', '原文／游戏内文本', '译文',
               '匹配状态', '条件／分支', '来源位置', '原始文本', '原始译文', '续段']
    widths = [10, 32, 42, 22, 20, 65, 65, 65, 16, 48, 55, 65, 65, 12]
    sheet = None
    sheet_n = 0
    row_count = 0
    try:
        for n, r in enumerate(records):
            if n % 500 == 0:
                check(f'正在导出 Excel · {n:,}/{len(records):,}')
            values = [r.seq, r.group, r.scene, r.category, r.speaker, r.display, r.original, r.translation,
                      r.status, r.context, r.source, r.raw, r.raw_translation]
            values = [str(v) for v in values]
            # Split by UTF-16 code units so astral characters cannot exceed Excel's limit.
            def chunks(s):
                parts, buf, units = [], [], 0
                for char in s:
                    size = 2 if ord(char) > 0xFFFF else 1
                    if units + size > 30000:
                        parts.append(''.join(buf)); buf = []; units = 0
                    buf.append(char); units += size
                parts.append(''.join(buf))
                return parts
            parts = [chunks(v) for v in values]
            count = max(map(len, parts))
            for part in range(count):
                if sheet is None or row_count >= 1_000_000:
                    if sheet:
                        sheet.auto_filter.ref = f'A1:N{row_count}'
                    sheet_n += 1
                    sheet = book.create_sheet('文本' if sheet_n == 1 else f'文本_{sheet_n}')
                    sheet.freeze_panes = 'F2'
                    for index, width in enumerate(widths):
                        sheet.column_dimensions[chr(65+index)].width = width
                    sheet.append([safe_cell(sheet, h, True) for h in headers])
                    row_count = 1
                row = [(p[part] if part < len(p) else (p[0] if i in (0, 1, 2, 3, 4, 8, 10) else ''))
                       for i, p in enumerate(parts)] + [f'{part+1}/{count}']
                sheet.append([safe_cell(sheet, value) for value in row])
                row_count += 1
        if sheet:
            sheet.auto_filter.ref = f'A1:N{row_count}'
        check('正在保存 Excel…')
        book.save(path)
    finally:
        # Close streaming XML writers even when the user cancels.
        for ws in book.worksheets:
            if not ws.closed:
                ws.close()
        book.close()


def export_records(result: Result, records: list[Record], target: str, fmt='txt', split=False,
                   bilingual=False, raw=False, progress=None, cancelled=None) -> Path:
    if not records:
        raise ValueError('当前范围没有可导出的文本。')
    if fmt not in ('txt', 'xlsx'):
        raise ValueError('不支持的导出格式。')
    def check(message):
        if cancelled and cancelled():
            raise InterruptedError('已取消导出。已生成的文件保存在本次导出目录中，目录内有未完成标记。')
        if progress:
            progress(message)
    check('准备导出…')
    folder = Path(target) / ('台词拾光_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
    folder.mkdir(parents=True, exist_ok=False)
    marker = folder / '导出未完成.txt'
    marker.write_text('导出未完成。请重新导出；.part 文件不能作为正式结果使用。', encoding='utf-8-sig')
    groups = defaultdict(list)
    if split:
        for r in records:
            groups[r.scene_id].append(r)
    else:
        groups['all'] = records
    files = []
    for index, subset in enumerate(groups.values(), 1):
        check(f'正在导出文件 {index}/{len(groups)}')
        label = (subset[0].group + '_' + subset[0].scene) if split else (result.title + '_汇总')
        name = f'{index:05d}_{safe_name(label)}.{fmt}'
        final = folder / name
        temporary = folder / (name + '.part')
        if fmt == 'txt':
            text_export(temporary, subset, result.title, bilingual, raw, check)
        else:
            excel_export(temporary, subset, result, check)
        temporary.replace(final)
        files.append({'file': name, 'records': len(subset), 'scene': subset[0].scene if split else '汇总'})
    (folder / '导出清单.json').write_text(json.dumps({'title': result.title, 'engine': result.engine,
        'source': result.root, 'translation_file': result.translation_file, 'records': len(records),
        'warnings': result.warnings, 'files': files}, ensure_ascii=False, indent=2), encoding='utf-8')
    marker.unlink()
    return folder
