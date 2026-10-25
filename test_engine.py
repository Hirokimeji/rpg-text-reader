import json
import tempfile
import unittest
from pathlib import Path

from engine import Translator, clean, extract, game_root, json_from_js, translation_candidates
from exporter import export_records


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        (self.root / 'data').mkdir()
        (self.root / 'js').mkdir()
        (self.root / 'js' / 'rmmz_core.js').write_text('// marker')
        self.write('System', {'gameTitle': '测试游戏', 'terms': {'messages': {'test': '=1+1'}}})
        self.write('MapInfos', [None, {'id': 1, 'name': '起点'}])
        cmds = [self.cmd(101, ['', 0, 0, 2, 'Alice']), self.cmd(401, ['Hello']), self.cmd(401, ['World']),
                self.cmd(102, [['Yes', 'No'], 0]), self.cmd(402, [0, 'Yes']),
                self.cmd(401, ['Partial'], 1), self.cmd(401, ['Missing'], 1), self.cmd(404, []),
                self.cmd(401, ['Hello']), self.cmd(355, ['const text = "safe";']), self.cmd(0, [])]
        pages = [{'conditions': {'switch1Valid': True, 'switch1Id': 2}, 'list': cmds},
                 {'conditions': {}, 'list': [self.cmd(401, ['Hello']), self.cmd(0, [])]}]
        self.write('Map001', {'displayName': 'Starting place', 'events': [None, {'id': 1, 'name': '对话', 'pages': pages}]})
        self.write('CommonEvents', [None, {'id': 1, 'name': '问候', 'list': [self.cmd(401, ['Hello'])]}])
        self.write('Troops', [None, {'id': 1, 'name': 'Slime', 'pages': [{'list': [self.cmd(401, ['Battle'])]}]}])
        self.translation = self.root / '翻译.json'
        self.translation.write_text(json.dumps({'Hello\nWorld': '你好\n世界', 'Hello': '你好', 'Partial': '部分',
            'Yes': '是', 'No': '否', 'Alice': '爱丽丝', 'Unused': '未关联', 'Blank': ''}), encoding='utf-8')

    def tearDown(self):
        self.tmp.cleanup()

    @staticmethod
    def cmd(code, parameters, indent=0):
        return {'code': code, 'parameters': parameters, 'indent': indent}

    def write(self, stem, data):
        (self.root / 'data' / (stem + '.json')).write_text(json.dumps(data), encoding='utf-8')

    def run_extract(self):
        return extract(str(self.root), str(self.translation))

    def test_paragraph_matching_and_scene_duplicates(self):
        result = self.run_extract()
        para = next(r for r in result.records if r.raw == 'Hello\nWorld')
        self.assertEqual(para.translation, '你好\n世界')
        self.assertEqual(para.speaker, '爱丽丝')
        self.assertIn('开关#2', para.context)
        partial = next(r for r in result.records if r.raw == 'Partial\nMissing')
        self.assertEqual(partial.status, '部分匹配')
        self.assertEqual(partial.translation, '部分\nMissing')
        self.assertIn('选择：是', partial.context)
        repeats = [r for r in result.records if r.raw == 'Hello' and r.category == '对白']
        self.assertEqual(len(repeats), 3)
        self.assertEqual(len({r.scene_id for r in repeats}), 3)
        self.assertTrue(any(r.raw == 'Unused' and r.group == '未关联场景' for r in result.records))
        self.assertTrue(any(r.category == '脚本／插件（待辨识）' for r in result.records))
        self.assertTrue(any(r.group == '战斗事件' for r in result.records))

    def test_no_translation_and_detection(self):
        result = extract(str(self.root))
        self.assertTrue(all(r.status == '游戏内文本' for r in result.records))
        self.assertEqual(game_root(self.root / 'data'), self.root.resolve())
        self.assertEqual(result.engine, 'RPG Maker MZ')

    def test_controls_preserve_unknown_and_variables(self):
        self.assertEqual(clean(r'\C[2]\N<小明>你好\V[3]\N[2]\!\Mystery[4]'), '你好〔变量#3〕〔角色#2〕\\Mystery[4]')

    def test_translation_discovery_validates_contents_and_skips_cache(self):
        for name, content in {
            'MTool_TilesetDefinedCache.json': '{"tile": "cache"}',
            'package.json': '{"name": "game"}',
            'empty.json': '{}', 'array.json': '["text"]',
            'nested.json': '{"settings": {"enabled": true}}',
            'number.json': '{"version": 1}', 'broken.json': '{',
        }.items():
            (self.root / name).write_text(content, encoding='utf-8')
        (self.root / 'binary.json').write_bytes(b'\xff\xfe\x00')
        folder = self.root / 'translations'
        folder.mkdir()
        other = folder / 'custom.json'
        other.write_text('{"Hello": ""}', encoding='utf-8-sig')
        self.assertEqual(set(translation_candidates(self.root)), {self.translation, other})
        with self.assertRaisesRegex(ValueError, '缓存或配置文件'):
            Translator(self.root / 'MTool_TilesetDefinedCache.json')

    def test_mv_translation_discovery_searches_www_and_game_directory(self):
        www = self.root / 'www'
        www.mkdir()
        nested = www / 'translation'
        nested.mkdir()
        other = nested / 'custom.json'
        other.write_text('{"Hello": "Hi"}', encoding='utf-8')
        self.assertEqual(set(translation_candidates(www)), {self.translation, other})

    def test_dictionary_only_empty_and_bad_formats(self):
        standalone = self.root / 'standalone'
        standalone.mkdir()
        target = standalone / 'dict.json'
        target.write_text(self.translation.read_text(), encoding='utf-8')
        result = extract(str(target))
        self.assertEqual(len(result.records), 8)
        self.assertTrue(all(r.group == '未关联场景' for r in result.records))
        self.assertTrue(any(r.status == '未匹配' for r in result.records))
        target.write_text('{"bad": {"nested": "value"}}')
        with self.assertRaisesRegex(ValueError, '字符串字典'):
            extract(str(target))

    def test_normalized_line_endings(self):
        translator = Translator(self.translation)
        self.assertEqual(translator.translate('Hello\r\nWorld'), ('你好\n世界', '已匹配'))
        self.assertEqual(translator.translate('\\N<Alice>\nHello')[1], '已匹配')

    def test_cancellation(self):
        with self.assertRaises(InterruptedError):
            extract(str(self.root), cancelled=lambda: True)

    def test_plugin_data_is_parsed_without_executing(self):
        literal = '[ /* comment */ {"name":"Demo","status":true,"parameters":{"text":"https://a/b/*x*/",},}, // end\n]; evil();'
        self.assertEqual(json_from_js(literal)[0]['parameters']['text'], 'https://a/b/*x*/')
        (self.root/'js'/'plugins.js').write_text('var $plugins = ' + literal, encoding='utf-8')
        self.assertTrue(any(r.group == '插件参数' for r in self.run_extract().records))

    def test_txt_excel_and_split(self):
        from openpyxl import load_workbook
        result = self.run_extract()
        before = {str(p): p.read_bytes() for p in (self.root / 'data').glob('*.json')}
        txt = export_records(result, result.records, str(self.root / 'out'), split=True, bilingual=True)
        self.assertEqual(len(list(txt.glob('*.txt'))), len({r.scene_id for r in result.records}))
        self.assertFalse((txt / '导出未完成.txt').exists())
        contents = '\n'.join(p.read_text(encoding='utf-8-sig') for p in txt.glob('*.txt'))
        self.assertIn('原文：Hello\nWorld', contents)
        self.assertIn('部分匹配', contents)
        excel = export_records(result, result.records, str(self.root / 'out'), fmt='xlsx')
        book = load_workbook(next(excel.glob('*.xlsx')), read_only=True)
        rows = list(book['文本'].iter_rows())
        self.assertEqual(len(rows)-1, len(result.records))
        formula_text = next(row[6] for row in rows[1:] if row[6].value == '=1+1')
        self.assertEqual(formula_text.data_type, 's')
        book.close()
        self.assertEqual(before, {str(p): p.read_bytes() for p in (self.root / 'data').glob('*.json')})

    def test_long_excel_text_is_not_lost(self):
        from openpyxl import load_workbook
        result = self.run_extract()
        record = result.records[0]
        record.original = '😀' * 17000 + '末尾'
        record.raw = record.original
        record.translation = ''
        output = export_records(result, [record], str(self.root / 'out'), fmt='xlsx')
        book = load_workbook(next(output.glob('*.xlsx')), read_only=True)
        data = list(book['文本'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(''.join(r[6] for r in data), record.original)
        self.assertEqual(len(data), 2)
        book.close()


if __name__ == '__main__':
    unittest.main()
